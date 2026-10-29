from openpyxl import load_workbook

def parse_playstation():

    workbook = load_workbook(
        "feeds/playstation_games_bad.xlsx",
        data_only=True
    )

    sheet = workbook["PlayStation Games"]

    games = []

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        if not any(value is not None for value in row):
            continue

        game = {
            "id": row[0],
            "name": row[1],
            "genre": row[2],
            "developer": row[3],
            "publisher": row[4],
            "release_date": row[5],
            "price": row[6],
            "currency": row[7],
            "age_rating": row[8],
            "user_rating": row[9],
            "platform": row[10],
            "image_url": row[11],
            "availability": row[12]
        }

        games.append(game)

    return games


if __name__ == "__main__":

    games = parse_playstation()

    print(f"Games found: {len(games)}")

    for game in games[:3]:
        print(game)